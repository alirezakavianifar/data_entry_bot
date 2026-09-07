import datetime
from pathlib import Path
import openpyxl
import pytest

from data.models import Client, RegistrationResult, RegistrationStatus
from data.excel_provider import ExcelDataProvider


def test_formatted_notes_on_failed_registration():
    """Verify that failed registrations on embargo sites output actionable error notes rather than betting cooling-off text."""
    fail_res = RegistrationResult(
        client_id="CLI_0001_test",
        client_name="Laura Francis",
        site_id="betfair",
        site_name="Betfair",
        status=RegistrationStatus.FAILED,
        email="lau.francis23@outlook.com",
        error_summary="Invalid phone number: We could not validate your mobile number, please check and try again.",
        notes="Phone validation failed on Betfair: 'We could not validate your mobile number, please check and try again.'. Please check client mobile number (7582070319)."
    )

    note = fail_res.formatted_notes
    assert "Please do not place any bets" not in note
    assert "We could not validate your mobile number" in note
    assert "7582070319" in note


def test_formatted_notes_on_paddypower_failed_registration():
    """Verify Paddy Power failed registration preserves failure note."""
    fail_res = RegistrationResult(
        client_id="CLI_0002_test",
        client_name="Laura Francis",
        site_id="paddypower",
        site_name="Paddy Power",
        status=RegistrationStatus.FAILED,
        email="lau.francis23@outlook.com",
        error_summary="Invalid phone number: We could not validate your mobile number",
        notes="Phone validation failed on Paddy Power: 'We could not validate your mobile number'"
    )

    note = fail_res.formatted_notes
    assert "Please do not place any bets" not in note
    assert "We could not validate your mobile number" in note


def test_formatted_notes_preserves_cooling_off_for_success():
    """Verify successful registrations on embargo sites still receive the required 25-day betting embargo statement."""
    succ_res = RegistrationResult(
        client_id="CLI_0003_test",
        client_name="Jake Fox",
        site_id="betfair",
        site_name="Betfair",
        status=RegistrationStatus.SUCCESS,
        email="jake.fox@example.com",
        account_reference="BF-REF-999"
    )

    note = succ_res.formatted_notes
    assert "successful. Please do not place any bets until within 25 days." in note
    assert "BF-REF-999" in note


def test_excel_record_failure_updates_client_details_notes(tmp_path):
    """Verify ExcelDataProvider updates the Notes column in Client Details tab on failure."""
    file_path = tmp_path / "Test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Client Details"

    # Setup headers
    ws.append(["Client", "First name", "Last name", "DOB", "Email", "Phone", "Address line 1", "Town / City", "Postcode", "Country", "Card used", "Allocated VA", "Notes"])
    ws.append(["Laura Francis", "Laura", "Francis", "23/04/1996", "lau.francis23@outlook.com", "07582070319", "10 High St", "London", "SW1A 1AA", "United Kingdom", "", "", ""])
    wb.save(file_path)

    provider = ExcelDataProvider(file_path=file_path)
    fail_res = RegistrationResult(
        client_id="CLI_0001_laura",
        client_name="Laura Francis",
        site_id="betfair",
        site_name="Betfair",
        status=RegistrationStatus.FAILED,
        email="lau.francis23@outlook.com",
        error_summary="Invalid phone number: We could not validate your mobile number, please check and try again.",
        notes="Phone number validation failed on Betfair: 'We could not validate your mobile number, please check and try again.'. Please check client mobile number (07582070319)."
    )

    success = provider.record_failure(fail_res)
    assert success is True

    # Inspect the saved workbook
    wb_after = openpyxl.load_workbook(file_path)
    sheet_after = wb_after["Client Details"]
    note_cell_val = sheet_after.cell(2, 13).value

    assert note_cell_val is not None
    assert "[Betfair]" in note_cell_val
    assert "We could not validate your mobile number" in note_cell_val
    assert "07582070319" in note_cell_val
