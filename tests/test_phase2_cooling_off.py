import datetime
from pathlib import Path
import openpyxl
import pytest

from data.models import RegistrationResult, RegistrationStatus
from data.excel_provider import ExcelDataProvider


def test_paddypower_embargo_note_generation():
    result = RegistrationResult(
        client_id="CLI_0001_test",
        client_name="Jake Fox",
        site_id="paddypower",
        site_name="Paddy Power",
        status=RegistrationStatus.SUCCESS,
        email="jake.fox@example.com",
        username="jakefox93",
        account_reference="PP-REF-12345"
    )

    note = result.formatted_notes
    assert "successful." in note
    assert "Please do not place any bets until within 25 days." in note
    assert "Signed up:" in note
    assert "Acceptable to bet on:" in note
    assert "PP-REF-12345" in note

    # Check date difference is exactly 25 days
    assert result.signup_date is not None
    assert result.acceptable_to_bet_date is not None
    assert (result.acceptable_to_bet_date - result.signup_date).days == 25


def test_betfair_embargo_note_generation():
    fixed_time = datetime.datetime(2026, 9, 6, 12, 0, 0)
    result = RegistrationResult(
        client_id="CLI_0002_test",
        client_name="Hannah Parkin",
        site_id="betfair",
        site_name="Betfair",
        status=RegistrationStatus.SUCCESS,
        email="hannah@example.com",
        username="hannahp",
        timestamp=fixed_time
    )

    note = result.formatted_notes
    assert "successful. Please do not place any bets until within 25 days." in note
    assert "Signed up: 06/09/2026" in note
    assert "Acceptable to bet on: 01/10/2026" in note


def test_standard_site_does_not_force_embargo():
    result = RegistrationResult(
        client_id="CLI_0003_test",
        client_name="Daniel Buckley",
        site_id="bettom",
        site_name="BetTOM",
        status=RegistrationStatus.SUCCESS,
        email="daniel@example.com",
        username="danb",
        account_reference="BT-REF-999"
    )

    note = result.formatted_notes
    assert note == "BT-REF-999"
    assert "Please do not place any bets" not in note


def test_excel_recording_embargo(tmp_path):
    wb_file = tmp_path / "test_sheet.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Client Details"
    wb.save(str(wb_file))

    provider = ExcelDataProvider(file_path=wb_file, output_sheet_name="Succesful Signuos")
    result = RegistrationResult(
        client_id="CLI_0004_test",
        client_name="Courtney Weaver",
        site_id="paddypower",
        site_name="Paddy Power",
        status=RegistrationStatus.SUCCESS,
        email="courtney@example.com",
        username="courtneyw"
    )

    success = provider.record_success(result)
    assert success is True

    # Read back
    wb_read = openpyxl.load_workbook(str(wb_file))
    ws = wb_read["Succesful Signuos"]
    assert ws.max_row == 2
    row2 = [ws.cell(2, c).value for c in range(1, 8)]
    assert row2[0] == "Courtney Weaver"
    assert row2[1] == "Paddy Power"
    assert row2[2] == "courtney@example.com"
    assert "successful. Please do not place any bets until within 25 days." in str(row2[6])
